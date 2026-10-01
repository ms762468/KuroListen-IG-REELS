"""Render editorially reviewed highlight TSV into Word files and append story manifest.
TSV columns: room number, pet, zero-based start/end message indices, title,
context, six spoken parts separated by ～, comma-separated evidence indices.
Run with the bundled Python runtime and python-docx. Review semantics before running;
render and inspect output before delivery. Exact source windows are idempotent.
"""
import csv,json,re,sys,pathlib,math,argparse
from collections import Counter
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from ingest import parse_csv
sys.stdout.reconfigure(encoding='utf-8')
p=argparse.ArgumentParser();p.add_argument('--project',type=pathlib.Path,required=True);p.add_argument('--specs',type=pathlib.Path,required=True);p.add_argument('--batch',required=True);p.add_argument('--raw-dir',default='data/raw/20261001');args=p.parse_args()
assert re.fullmatch(r'[A-Za-z0-9_-]+',args.batch)
repo=args.project.resolve();rawdir=repo/args.raw_dir
out=repo/'outputs'/args.batch;out.mkdir(parents=True,exist_ok=True)
manifestpath=repo/'data/story_manifest.json';state=json.loads(manifestpath.read_text(encoding='utf-8'));existing=state['stories']
chatids={str(s['source_room_number']):s['chat_id'] for s in existing}
counts=Counter()
for s in existing:counts[str(s['source_room_number'])]=max(counts[str(s['source_room_number'])],int(s['story_id'].split('-')[1]))
specs=[r for r in csv.reader(args.specs.open(encoding='utf-8'),delimiter='\t') if r]
manifest=[]
def safe(s):return re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s).strip()[:65]
def font(run,size=11,bold=False,color='000000'):
 run.font.name='Microsoft JhengHei';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
 rf=run._element.get_or_add_rPr().get_or_add_rFonts()
 for attr in list(rf.attrib):
  if 'Theme' in attr:del rf.attrib[attr]
 rf.set(qn('w:eastAsia'),'Microsoft JhengHei')
def para(doc,text,size=10.5,bold=False,after=4):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after);font(p.add_run(text),size,bold);return p
for r in specs:
 rid,pet,start,end,title,context,spoken,evidence_indexes=r;start=int(start);end=int(end)
 src=next(rawdir.glob(rid+'_*.csv'));parsed=parse_csv(src.read_bytes(),src.name)
 if any(s['source_room_number']==int(rid) and s['source_records']==[start+5,end+5] for s in existing):continue
 selected=parsed['messages'][start:end+1];keys=parsed['message_keys'][start:end+1]
 user=next((m['sender'] for m in parsed['messages'] if m['role']=='User'),'未提供')
 user=re.sub(r'\d{8,}', '[電話略]',user)
 dates=list(dict.fromkeys(m['date'] for m in selected)); date='、'.join(dates)
 parts=spoken.split('～')
 if len(parts)==5:
  ix=max(range(5),key=lambda i:len(parts[i]));s=parts[ix]
  boundaries=[m.end() for m in re.finditer('[。？！；，]',s) if 5<m.end()<len(s)-5]
  split=min(boundaries,key=lambda j:abs(j-len(s)/2))
  parts[ix:ix+1]=[s[:split],s[split:]]
 assert len(parts)==6,(rid,len(parts))
 charcounts=[len(re.findall(r'[\u4e00-\u9fffA-Za-z0-9]',s)) for s in parts]
 seconds=[n/4.5+0.5 for n in charcounts]; total=math.ceil(sum(seconds))
 assert total<=30,(rid,total)
 counts[rid]+=1; sid=f'R{int(rid):03d}-{counts[rid]:02d}'
 filename=f'{sid}_{safe(title)}.docx'
 evidence=[parsed['messages'][int(i)] for i in evidence_indexes.split(',')]
 confirmations=[m for m in evidence if m['role']=='User']
 assert confirmations, (rid,'Missing caregiver confirmation')
 keywords=[re.sub(r'\s+',' ',m['text'])[:32] for m in evidence if m['role']=='Account'][:1]
 keywords += [re.sub(r'\s+',' ',confirmations[0]['text'])[:32]]
 confirmation=' ／ '.join(re.sub(r'\s+',' ',m['text'])[:45] for m in confirmations[:2])
 d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
 sec.top_margin=Cm(1.45);sec.bottom_margin=Cm(1.45);sec.left_margin=Cm(1.65);sec.right_margin=Cm(1.65)
 normal=d.styles['Normal'];normal.font.name='Microsoft JhengHei';normal.font.size=Pt(10.5)
 normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft JhengHei')
 normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(4)
 for sn in ('Title','Heading 1','Heading 2'):
  d.styles[sn].font.color.rgb=RGBColor(0,0,0);d.styles[sn].font.name='Microsoft JhengHei'
 p=d.add_paragraph(style='Title');p.paragraph_format.space_after=Pt(5);font(p.add_run(title),19,True)
 para(d,f'{sid}  ·  KuroListen  ·  本人口述與照片  ·  約 {total} 秒',9.5,False,8)
 para(d,'前情提要  僅供回查',11,True)
 para(d,f'LINE 帳號：{user}　　毛孩：{pet}',10)
 para(d,f'對話日期：{date}（沿用匯出日期；來源時區 {parsed["timezone"]}）',9)
 para(d,context,10)
 para(d,'家長原文確認：'+confirmation,9)
 para(d,'LINE 搜尋：'+' ／ '.join(keywords),9)
 para(d,f'來源：{src.name}；CSV 第 {start+5}–{end+5} 筆邏輯列（含前三列資料及標題列，非換行行數）。',8.5,False,7)
 para(d,'口說稿與畫面',11,True)
 table=d.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
 widths=[1.45,10.65,5.55]
 for cell,w in zip(table.rows[0].cells,widths):cell.width=Cm(w)
 for cell,text in zip(table.rows[0].cells,['秒數','口說台詞','照片與字幕']):
  cell.text=text
  shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF0');cell._tc.get_or_add_tcPr().append(shading)
  for run in cell.paragraphs[0].runs:font(run,9.5,True)
 cues=[('本人看鏡頭','封面主題'),('毛孩一般肖像','故事背景'),('同張照片慢推近','引用短句'),('切回本人或照片','笑點停半秒'),('本人自然收尾','一個思考觀念'),('毛孩照片與帳號','留言問句')]
 elapsed=0
 for i,(part,duration) in enumerate(zip(parts,seconds)):
  row=table.add_row();stop=round(elapsed+duration) if i<5 else total
  texts=[f'{round(elapsed)}–{stop}',part,cues[i][0]+'\n'+cues[i][1]]
  for ci,(cell,w,text) in enumerate(zip(row.cells,widths,texts)):
   cell.width=Cm(w);cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    for run in p.runs:font(run,11 if ci==1 else 9.5)
  elapsed+=duration
 for row in table.rows:
  trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   tc=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tc.append(borders)
 para(d,'拍攝及改編',11,True,4)
 para(d,'以敘述為主，引用時可稍微變聲；不需演戲。照片尚未包含於 CSV，請用獲准肖像；沒有情境實拍照時，用一般照片搭配文字示意。',9)
 para(d,'畫面小字：溝通經驗分享・對話經濃縮改編。原對話支持的內容見前情提要；比喻與敘述者吐槽為新寫，不作動物逐字原話或效果保證。',9)
 para(d,'封面：'+title+'\n貼文：'+parts[4]+parts[5],9)
 para(d,f'口說約 {sum(charcounts)} 字，以每秒約 4.5 字加停頓估算；未經真人錄音實測。前情提要及姓名不放入公開字幕。',8.5)
 d.core_properties.title=title;d.core_properties.author='KuroListen';d.core_properties.subject='Reels 口述故事腳本'
 # Word 2010 can inherit a blue paragraph rule from its default Title style.
 for root in (d._element,d.styles._element):
  for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.save(out/filename)
 manifest.append({'story_id':sid,'chat_id':chatids[rid],'source_room_number':int(rid),'account_name':user,'pet':pet,'dates':dates,'source_csv':str((rawdir/src.name).relative_to(repo)).replace('\\','/'),'source_records':[start+5,end+5],'source_message_keys':keys,'scenario_summary':context,'title':title,'keywords':keywords,'spoken_parts':parts,'estimated_seconds':total,'batch':args.batch,'selection_reason':'家長具體確認的新情境','status':'generated','output':str((out/filename).relative_to(repo)).replace('\\','/'),'evidence':[{'role':m['role'],'record':m['record'],'text':m['text']} for m in evidence]})

state['stories']=existing+manifest
manifestpath.write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
batchstories=[s for s in state['stories'] if s.get('batch')==args.batch]
with (out/'腳本索引.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['編號','LINE帳號','毛孩','日期','主題','估計秒數','Word檔'])
 for s in batchstories:w.writerow([s['story_id'],s['account_name'],s['pet'],'、'.join(s['dates']),s['title'],s['estimated_seconds'],pathlib.Path(s['output']).name])
(out/'使用說明.txt').write_text(f'新增 {len(batchstories)} 篇家長確認的生活故事，每篇一個 Word，沿用口述加照片，不需演戲。\n內含家長原文確認、LINE 搜尋資訊及來源列，姓名只供內部回查。\n口述時間為文字估算，正式拍攝請試讀；照片不在 CSV 中，需自行選用獲准素材。\n已知資訊、家長更正與照片出現順序不可改寫為事前全部說中。\n',encoding='utf-8')
print(json.dumps({'new':len(manifest),'batch':len(batchstories),'total':len(state['stories']),'seconds':[min(s['estimated_seconds'] for s in batchstories),max(s['estimated_seconds'] for s in batchstories)]}))
